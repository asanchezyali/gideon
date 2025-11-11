"""
Ejemplo: Soporte Multi-Formato
================================

Este ejemplo muestra cómo extender Gideon para soportar múltiples formatos
de documentos más allá de PDF.

Formatos soportados:
- PDF (actual)
- Word (.docx)
- PowerPoint (.pptx)
- EPUB (libros electrónicos)
- Text (.txt)
- Markdown (.md)
- HTML
- Imágenes con OCR (.jpg, .png, .tiff)
"""

from pathlib import Path
from typing import Dict, Any, Optional, List
from abc import ABC, abstractmethod
from dataclasses import dataclass


@dataclass
class ExtractedDocument:
    """Extracted document content and metadata."""
    content: str
    metadata: Dict[str, Any]
    format: str
    extraction_quality: float  # 0.0 - 1.0


class BaseExtractor(ABC):
    """Base class for document extractors."""

    @abstractmethod
    async def extract_content(self, file_path: Path) -> str:
        """Extract text content from document."""
        pass

    @abstractmethod
    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract metadata (title, author, date, etc.)."""
        pass

    async def extract(self, file_path: Path) -> ExtractedDocument:
        """Extract both content and metadata."""
        content = await self.extract_content(file_path)
        metadata = await self.extract_metadata(file_path)

        return ExtractedDocument(
            content=content,
            metadata=metadata,
            format=file_path.suffix,
            extraction_quality=self._estimate_quality(content, metadata)
        )

    def _estimate_quality(self, content: str, metadata: Dict) -> float:
        """Estimate extraction quality score."""
        score = 0.5  # Base score

        # Content quality
        if len(content) > 100:
            score += 0.2
        if len(content) > 1000:
            score += 0.1

        # Metadata quality
        if metadata.get("title"):
            score += 0.1
        if metadata.get("author"):
            score += 0.1

        return min(score, 1.0)


# ============================================================================
# PDF Extractor (existing implementation)
# ============================================================================

class PDFExtractor(BaseExtractor):
    """Extract content from PDF files."""

    async def extract_content(self, file_path: Path) -> str:
        """Extract text from PDF."""
        from PyPDF2 import PdfReader

        try:
            reader = PdfReader(str(file_path))
            pages = reader.pages[:5]  # First 5 pages
            return "\n".join([page.extract_text() for page in pages if page.extract_text()])
        except Exception as e:
            print(f"Error extracting PDF: {e}")
            return ""

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract PDF metadata."""
        from PyPDF2 import PdfReader

        try:
            reader = PdfReader(str(file_path))
            metadata = reader.metadata or {}

            return {
                "title": metadata.get("/Title", ""),
                "author": metadata.get("/Author", ""),
                "subject": metadata.get("/Subject", ""),
                "creator": metadata.get("/Creator", ""),
                "producer": metadata.get("/Producer", ""),
                "creation_date": metadata.get("/CreationDate", ""),
                "pages": len(reader.pages)
            }
        except Exception as e:
            print(f"Error extracting PDF metadata: {e}")
            return {}


# ============================================================================
# Word Document Extractor (.docx)
# ============================================================================

class WordExtractor(BaseExtractor):
    """Extract content from Word documents (.docx)."""

    async def extract_content(self, file_path: Path) -> str:
        """Extract text from Word document."""
        try:
            from docx import Document

            doc = Document(str(file_path))

            # Extract all paragraphs
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

            # Extract text from tables
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    table_text.append(" | ".join(row_text))

            content = "\n".join(paragraphs)
            if table_text:
                content += "\n\nTables:\n" + "\n".join(table_text)

            return content

        except Exception as e:
            print(f"Error extracting Word document: {e}")
            return ""

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract Word document metadata."""
        try:
            from docx import Document

            doc = Document(str(file_path))
            core_props = doc.core_properties

            return {
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "keywords": core_props.keywords or "",
                "created": core_props.created,
                "modified": core_props.modified,
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables)
            }
        except Exception as e:
            print(f"Error extracting Word metadata: {e}")
            return {}


# ============================================================================
# PowerPoint Extractor (.pptx)
# ============================================================================

class PowerPointExtractor(BaseExtractor):
    """Extract content from PowerPoint presentations (.pptx)."""

    async def extract_content(self, file_path: Path) -> str:
        """Extract text from PowerPoint."""
        try:
            from pptx import Presentation

            prs = Presentation(str(file_path))

            content_parts = []

            for i, slide in enumerate(prs.slides, 1):
                slide_content = [f"=== Slide {i} ==="]

                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_content.append(shape.text)

                content_parts.append("\n".join(slide_content))

            return "\n\n".join(content_parts)

        except Exception as e:
            print(f"Error extracting PowerPoint: {e}")
            return ""

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract PowerPoint metadata."""
        try:
            from pptx import Presentation

            prs = Presentation(str(file_path))
            core_props = prs.core_properties

            return {
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "keywords": core_props.keywords or "",
                "created": core_props.created,
                "modified": core_props.modified,
                "slides": len(prs.slides)
            }
        except Exception as e:
            print(f"Error extracting PowerPoint metadata: {e}")
            return {}


# ============================================================================
# EPUB Extractor (eBooks)
# ============================================================================

class EPUBExtractor(BaseExtractor):
    """Extract content from EPUB files."""

    async def extract_content(self, file_path: Path) -> str:
        """Extract text from EPUB."""
        try:
            import ebooklib
            from ebooklib import epub
            from bs4 import BeautifulSoup

            book = epub.read_epub(str(file_path))

            content_parts = []

            for item in book.get_items():
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    soup = BeautifulSoup(item.get_content(), 'html.parser')
                    text = soup.get_text()
                    if text.strip():
                        content_parts.append(text)

            return "\n\n".join(content_parts)

        except Exception as e:
            print(f"Error extracting EPUB: {e}")
            return ""

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract EPUB metadata."""
        try:
            import ebooklib
            from ebooklib import epub

            book = epub.read_epub(str(file_path))

            return {
                "title": book.get_metadata('DC', 'title')[0][0] if book.get_metadata('DC', 'title') else "",
                "author": book.get_metadata('DC', 'creator')[0][0] if book.get_metadata('DC', 'creator') else "",
                "language": book.get_metadata('DC', 'language')[0][0] if book.get_metadata('DC', 'language') else "",
                "publisher": book.get_metadata('DC', 'publisher')[0][0] if book.get_metadata('DC', 'publisher') else "",
                "date": book.get_metadata('DC', 'date')[0][0] if book.get_metadata('DC', 'date') else "",
            }
        except Exception as e:
            print(f"Error extracting EPUB metadata: {e}")
            return {}


# ============================================================================
# Text File Extractor (.txt, .md)
# ============================================================================

class TextExtractor(BaseExtractor):
    """Extract content from plain text files."""

    async def extract_content(self, file_path: Path) -> str:
        """Read text file."""
        try:
            return file_path.read_text(encoding='utf-8')
        except UnicodeDecodeError:
            # Try other encodings
            for encoding in ['latin-1', 'cp1252', 'iso-8859-1']:
                try:
                    return file_path.read_text(encoding=encoding)
                except UnicodeDecodeError:
                    continue
            return ""

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract basic file metadata."""
        stat = file_path.stat()
        return {
            "size": stat.st_size,
            "created": stat.st_ctime,
            "modified": stat.st_mtime,
            "lines": len(file_path.read_text().splitlines())
        }


# ============================================================================
# HTML Extractor
# ============================================================================

class HTMLExtractor(BaseExtractor):
    """Extract content from HTML files."""

    async def extract_content(self, file_path: Path) -> str:
        """Extract text from HTML."""
        try:
            from bs4 import BeautifulSoup

            html_content = file_path.read_text(encoding='utf-8')
            soup = BeautifulSoup(html_content, 'html.parser')

            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()

            text = soup.get_text()

            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)

            return text

        except Exception as e:
            print(f"Error extracting HTML: {e}")
            return ""

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract HTML metadata."""
        try:
            from bs4 import BeautifulSoup

            html_content = file_path.read_text(encoding='utf-8')
            soup = BeautifulSoup(html_content, 'html.parser')

            return {
                "title": soup.title.string if soup.title else "",
                "meta_description": soup.find("meta", {"name": "description"})["content"]
                    if soup.find("meta", {"name": "description"}) else "",
                "meta_keywords": soup.find("meta", {"name": "keywords"})["content"]
                    if soup.find("meta", {"name": "keywords"}) else "",
                "meta_author": soup.find("meta", {"name": "author"})["content"]
                    if soup.find("meta", {"name": "author"}) else "",
            }
        except Exception as e:
            print(f"Error extracting HTML metadata: {e}")
            return {}


# ============================================================================
# Image OCR Extractor (.jpg, .png, .tiff)
# ============================================================================

class ImageOCRExtractor(BaseExtractor):
    """Extract text from images using OCR."""

    async def extract_content(self, file_path: Path) -> str:
        """Extract text from image using OCR."""
        try:
            from PIL import Image
            import pytesseract

            # Open image
            image = Image.open(file_path)

            # Perform OCR
            text = pytesseract.image_to_string(image)

            return text

        except Exception as e:
            print(f"Error extracting text from image: {e}")
            return ""

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Extract image metadata."""
        try:
            from PIL import Image

            image = Image.open(file_path)

            return {
                "format": image.format,
                "mode": image.mode,
                "size": image.size,
                "width": image.width,
                "height": image.height,
                "exif": dict(image.getexif()) if hasattr(image, 'getexif') else {}
            }
        except Exception as e:
            print(f"Error extracting image metadata: {e}")
            return {}


# ============================================================================
# Extractor Factory
# ============================================================================

class ExtractorFactory:
    """Factory for creating appropriate extractors based on file type."""

    EXTRACTOR_MAP = {
        ".pdf": PDFExtractor,
        ".docx": WordExtractor,
        ".pptx": PowerPointExtractor,
        ".epub": EPUBExtractor,
        ".txt": TextExtractor,
        ".md": TextExtractor,
        ".html": HTMLExtractor,
        ".htm": HTMLExtractor,
        ".jpg": ImageOCRExtractor,
        ".jpeg": ImageOCRExtractor,
        ".png": ImageOCRExtractor,
        ".tiff": ImageOCRExtractor,
        ".tif": ImageOCRExtractor,
    }

    @classmethod
    def get_extractor(cls, file_path: Path) -> BaseExtractor:
        """
        Get appropriate extractor for file.

        Args:
            file_path: Path to file

        Returns:
            Extractor instance

        Raises:
            ValueError: If file format not supported
        """
        extension = file_path.suffix.lower()

        if extension not in cls.EXTRACTOR_MAP:
            supported = ", ".join(cls.EXTRACTOR_MAP.keys())
            raise ValueError(
                f"Unsupported file format: {extension}\n"
                f"Supported formats: {supported}"
            )

        extractor_class = cls.EXTRACTOR_MAP[extension]
        return extractor_class()

    @classmethod
    def get_supported_extensions(cls) -> List[str]:
        """Get list of supported file extensions."""
        return list(cls.EXTRACTOR_MAP.keys())

    @classmethod
    def is_supported(cls, file_path: Path) -> bool:
        """Check if file format is supported."""
        return file_path.suffix.lower() in cls.EXTRACTOR_MAP


# ============================================================================
# Enhanced FileService
# ============================================================================

class EnhancedFileService:
    """
    Enhanced file service with multi-format support.

    This would replace/extend the current FileService.
    """

    @staticmethod
    async def extract_content(file_path: Path) -> str:
        """Extract content from any supported file format."""
        try:
            extractor = ExtractorFactory.get_extractor(file_path)
            return await extractor.extract_content(file_path)
        except Exception as e:
            print(f"Error extracting content from {file_path}: {e}")
            return ""

    @staticmethod
    async def extract_document(file_path: Path) -> ExtractedDocument:
        """Extract full document (content + metadata)."""
        extractor = ExtractorFactory.get_extractor(file_path)
        return await extractor.extract(file_path)

    @staticmethod
    def get_files_by_extensions(
        directory: Path,
        extensions: Optional[List[str]] = None
    ) -> List[Path]:
        """
        Get all files with specified extensions.

        Args:
            directory: Directory to search
            extensions: List of extensions (e.g., [".pdf", ".docx"])
                       If None, use all supported extensions

        Returns:
            List of matching file paths
        """
        if extensions is None:
            extensions = ExtractorFactory.get_supported_extensions()

        files = []
        for ext in extensions:
            files.extend(directory.rglob(f"*{ext}"))

        return files


# ============================================================================
# CLI Command Example
# ============================================================================

async def process_mixed_formats(directory: Path):
    """Example: Process directory with mixed file formats."""
    service = EnhancedFileService()

    # Get all supported files
    files = service.get_files_by_extensions(directory)

    print(f"Found {len(files)} files\n")

    # Group by format
    by_format = {}
    for file in files:
        ext = file.suffix.lower()
        if ext not in by_format:
            by_format[ext] = []
        by_format[ext].append(file)

    print("Files by format:")
    for ext, file_list in by_format.items():
        print(f"  {ext}: {len(file_list)} files")

    print("\nProcessing files...\n")

    # Extract content from all files
    for file in files[:5]:  # First 5 as example
        try:
            doc = await service.extract_document(file)

            print(f"File: {file.name}")
            print(f"Format: {doc.format}")
            print(f"Quality: {doc.extraction_quality:.2%}")
            print(f"Content length: {len(doc.content)} chars")
            print(f"Metadata: {list(doc.metadata.keys())}")
            print("-" * 60)

        except Exception as e:
            print(f"Error processing {file.name}: {e}")


# Example usage
if __name__ == "__main__":
    import asyncio

    # Test with mixed formats
    # asyncio.run(process_mixed_formats(Path("./documents")))

    # Check supported formats
    print("Supported formats:")
    for ext in ExtractorFactory.get_supported_extensions():
        print(f"  - {ext}")
