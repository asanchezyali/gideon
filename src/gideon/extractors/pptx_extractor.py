"""PowerPoint (.pptx) extractor."""

from pathlib import Path
from typing import Dict, Any
from pptx import Presentation

from .base import BaseExtractor


class PptxExtractor(BaseExtractor):
    """Extractor for PowerPoint presentations (.pptx)."""

    async def extract_content(self, file_path: Path) -> str:
        """
        Extract text content from PowerPoint presentation.

        Args:
            file_path: Path to .pptx file

        Returns:
            Extracted text content

        Raises:
            FileNotFoundError: If file doesn't exist
            ValueError: If file cannot be read
        """
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        try:
            prs = Presentation(str(file_path))
            slides_text = []

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_content = [f"=== Slide {slide_num} ==="]

                # Extract text from all shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_content.append(shape.text.strip())

                    # Extract text from tables
                    if shape.shape_type == 19:  # Table
                        try:
                            for row in shape.table.rows:
                                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                                if row_text:
                                    slide_content.append(" | ".join(row_text))
                        except:
                            pass

                if len(slide_content) > 1:  # Has content beyond header
                    slides_text.append("\n".join(slide_content))

            content = "\n\n".join(slides_text)

            if not content.strip():
                return "Empty presentation"

            return content

        except Exception as e:
            raise ValueError(f"Failed to extract content from {file_path}: {str(e)}")

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """
        Extract metadata from PowerPoint presentation.

        Args:
            file_path: Path to .pptx file

        Returns:
            Dictionary containing metadata

        Raises:
            FileNotFoundError: If file doesn't exist
            ValueError: If metadata cannot be extracted
        """
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        try:
            prs = Presentation(str(file_path))
            core_props = prs.core_properties

            metadata = {
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "keywords": core_props.keywords or "",
                "created": str(core_props.created) if core_props.created else "",
                "modified": str(core_props.modified) if core_props.modified else "",
                "last_modified_by": core_props.last_modified_by or "",
                "revision": core_props.revision or 0,
                "num_slides": len(prs.slides),
                "slide_width": prs.slide_width,
                "slide_height": prs.slide_height,
                "file_format": "pptx",
            }

            return metadata

        except Exception as e:
            raise ValueError(f"Failed to extract metadata from {file_path}: {str(e)}")
