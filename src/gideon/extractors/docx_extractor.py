"""Word document (.docx) extractor."""

from pathlib import Path
from typing import Dict, Any
import docx

from .base import BaseExtractor


class DocxExtractor(BaseExtractor):
    """Extractor for Word documents (.docx)."""

    async def extract_content(self, file_path: Path) -> str:
        """
        Extract text content from Word document.

        Args:
            file_path: Path to .docx file

        Returns:
            Extracted text content

        Raises:
            FileNotFoundError: If file doesn't exist
            ValueError: If file cannot be read
        """
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        try:
            doc = docx.Document(str(file_path))

            # Extract text from all paragraphs
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

            # Extract text from tables
            table_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        table_text.append(" | ".join(row_text))

            # Combine all text
            all_text = paragraphs + table_text
            content = "\n\n".join(all_text)

            if not content.strip():
                return "Empty document"

            return content

        except Exception as e:
            raise ValueError(f"Failed to extract content from {file_path}: {str(e)}")

    async def extract_metadata(self, file_path: Path) -> Dict[str, Any]:
        """
        Extract metadata from Word document.

        Args:
            file_path: Path to .docx file

        Returns:
            Dictionary containing metadata

        Raises:
            FileNotFoundError: If file doesn't exist
            ValueError: If metadata cannot be extracted
        """
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        try:
            doc = docx.Document(str(file_path))
            core_props = doc.core_properties

            metadata = {
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "keywords": core_props.keywords or "",
                "created": str(core_props.created) if core_props.created else "",
                "modified": str(core_props.modified) if core_props.modified else "",
                "last_modified_by": core_props.last_modified_by or "",
                "revision": core_props.revision or 0,
                "num_paragraphs": len(doc.paragraphs),
                "num_tables": len(doc.tables),
                "num_sections": len(doc.sections),
                "file_format": "docx",
            }

            return metadata

        except Exception as e:
            raise ValueError(f"Failed to extract metadata from {file_path}: {str(e)}")
